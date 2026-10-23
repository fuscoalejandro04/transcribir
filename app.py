import streamlit as st
import re
import io
from docx import Document
from youtube_transcript_api import YouTubeTranscriptApi
from youtube_transcript_api._errors import (
    TranscriptsDisabled,
    NoTranscriptFound,
    VideoUnavailable,
    CouldNotRetrieveTranscript,
)

# Configuración de la página (opcional, pero mejora la apariencia)
st.set_page_config(page_title="YouTube a Word", page_icon="📄", layout="centered")

st.title("📄 Transcripción de YouTube a Word")
st.caption("Pega la URL de un video de YouTube para descargar su transcripción en formato .docx")

def extraer_id_video(url: str) -> str | None:
    """Extrae el ID de 11 caracteres de una URL de YouTube."""
    patrones = [
        r'(?:youtube\.com\/watch\?v=)([a-zA-Z0-9_-]{11})',
        r'(?:youtu\.be\/)([a-zA-Z0-9_-]{11})',
        r'(?:youtube\.com\/embed\/)([a-zA-Z0-9_-]{11})',
        r'(?:youtube\.com\/shorts\/)([a-zA-Z0-9_-]{11})',
    ]
    for patron in patrones:
        match = re.search(patron, url)
        if match:
            return match.group(1)
    return None

def generar_docx(texto: str, url_video: str) -> bytes:
    """Crea un documento de Word en memoria y devuelve sus bytes."""
    doc = Document()
    doc.add_heading("Transcripción del Video de YouTube", 0)
    doc.add_paragraph(f"URL original: {url_video}")
    doc.add_paragraph("")  # Línea en blanco de separación
    
    # Dividir en párrafos para mejor legibilidad
    parrafos = re.split(r'(?<=[.!?]) +', texto)
    for parrafo in parrafos:
        if parrafo.strip():
            doc.add_paragraph(parrafo.strip())
    
    # Guardar en un buffer de memoria
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def obtener_transcripcion(video_id: str) -> str | None:
    """Intenta obtener la transcripción, priorizando español e inglés."""
    try:
        lista_transcripciones = YouTubeTranscriptApi.list_transcripts(video_id)
        
        # Estrategia de selección de idioma
        transcript = None
        for idioma in ['es', 'en']:
            try:
                # Intenta manual primero, luego generada
                try:
                    transcript = lista_transcripciones.find_manually_created_transcript([idioma])
                except NoTranscriptFound:
                    transcript = lista_transcripciones.find_generated_transcript([idioma])
                break  # Si encuentra una, sale del bucle
            except NoTranscriptFound:
                continue  # Prueba el siguiente idioma
        
        if transcript is None:
            # Si no hay es/en, toma la primera disponible
            transcript = list(lista_transcripciones)[0]
        
        datos = transcript.fetch()
        texto = " ".join([item['text'] for item in datos])
        return texto
        
    except TranscriptsDisabled:
        st.error("❌ El video tiene los subtítulos desactivados o no dispone de ellos.")
    except NoTranscriptFound:
        st.error("❌ No se encontró ninguna transcripción para este video.")
    except VideoUnavailable:
        st.error("❌ El video no está disponible (puede ser privado o haber sido eliminado).")
    except CouldNotRetrieveTranscript as e:
        st.error(f"❌ No se pudo recuperar la transcripción. Error: {e}")
    except Exception as e:
        st.error(f"❌ Ocurrió un error inesperado: {e}")
    
    return None

# --- Interfaz de Usuario ---
with st.form("url_form"):
    url_input = st.text_input(
        "🔗 URL del video de YouTube",
        placeholder="https://www.youtube.com/watch?v=...",
        help="También funciona con enlaces cortos (youtu.be) o Shorts."
    )
    submitted = st.form_submit_button("Obtener transcripción")

if submitted and url_input:
    video_id = extraer_id_video(url_input)
    
    if not video_id:
        st.warning("⚠️ No se pudo identificar un ID de video válido en la URL. Verifícala.")
    else:
        with st.spinner("Buscando y procesando la transcripción..."):
            texto_transcripcion = obtener_transcripcion(video_id)
        
        if texto_transcripcion:
            st.success("✅ ¡Transcripción obtenida con éxito!")
            
            # Generar el docx
            docx_bytes = generar_docx(texto_transcripcion, url_input)
            
            # Botón de descarga (el MIME type correcto para docx)
            st.download_button(
                label="📥 Descargar archivo Word (.docx)",
                data=docx_bytes,
                file_name=f"transcripcion_{video_id}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
            )
            
            # Opcional: mostrar una vista previa
            with st.expander("👀 Vista previa del texto"):
                st.write(texto_transcripcion[:1000] + "..." if len(texto_transcripcion) > 1000 else texto_transcripcion)

elif submitted and not url_input:
    st.warning("⚠️ Por favor, ingresa una URL.")
